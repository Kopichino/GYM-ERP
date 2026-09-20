"""Every gym-owned resource, tried from the wrong gym and by the wrong person.

Two unrelated gyms, each with an admin, a trainer and a member, plus a second
member at gym A. Gym B's records are the targets.

For each resource, gym A's people try to read, change and delete gym B's row --
through their own gym's URL and through gym B's -- and try to reach it through
the `?member=` style filters the lists accept. Every attempt must be refused
with a 403 or 404 (or a 405 where the route has no such method) and must leave
the row exactly as it was. Every target is also read once by the gym B person
entitled to it, so a refusal can never pass merely because an id was wrong.

The same-gym half covers a member reaching another member's rows and a trainer
reaching a member who is not theirs. Malformed and unknown ids must never be a
server error.
"""

import io
from unittest.mock import patch
from datetime import time, timedelta
from decimal import Decimal

import openpyxl
from django.core import mail
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import override_settings
from django.utils import timezone

from accounts.models import Role, User
from attendance.models import CheckInOut
from billing.models import Payment, PaymentMethod, PaymentOrder, Plan
from billing.services import record_payment
from bodystats.models import BodyMeasurement, MemberGoal
from crm.models import Enquiry
from expenses.models import Expense, ExpenseCategory
from gallery.models import GalleryPost
from invoicing.services import issue_invoice
from notifications.services import EXPIRY_WINDOWS, expiring_members
from nutrition.models import DietDay, DietMeal, DietPlan
from pt.models import PTSession
from referrals.models import Referral
from reports.models import SavedReport
from schedule_app.models import ClassBooking, ClassSession
from tenancy import context
from workouts.models import Exercise, SplitDay, WorkoutLog, WorkoutSession, WorkoutSplit

from .testing import TwoGymsTestCase, make_person, png_bytes

#: What a refusal may look like. 405 only means the route has no such method,
#: which is decided before any object is looked up, so it reveals nothing.
REFUSED = {403, 404}
REFUSED_WRITE = {403, 404, 405}
#: A stranger writing into a gym they do not belong to.
REFUSED_STRANGER = {401, 403, 404}

#: (label, path template, who at the target gym may read it)
OBJECT_ROUTES = [
    ("member account", "/auth/admin/users/{member}/", "admin"),
    ("trainer's member", "/auth/trainer/members/{member}/", "trainer"),
    ("payment", "/billing/admin/payments/{payment}/", "admin"),
    ("plan", "/billing/plans/{plan}/", "admin"),
    ("invoice", "/invoices/{invoice}/", "member"),
    ("invoice pdf", "/invoices/{invoice}/pdf/", "member"),
    ("workout session", "/workouts/sessions/{session}/", "member"),
    ("workout log", "/workouts/logs/{log}/", "member"),
    ("split", "/workouts/splits/{split}/", "member"),
    ("split day", "/workouts/split-days/{split_day}/", "member"),
    ("diet plan", "/nutrition/plans/{diet}/", "member"),
    ("diet day", "/nutrition/days/{diet_day}/", "member"),
    ("diet meal", "/nutrition/meals/{meal}/", "member"),
    ("measurement", "/bodystats/measurements/{measurement}/", "member"),
    ("goal", "/bodystats/goals/{goal}/", "member"),
    ("pt session", "/pt/sessions/{pt}/", "member"),
    ("class", "/schedule/{klass}/", "member"),
    ("class roster", "/schedule/{klass}/roster/", "trainer"),
    ("saved report", "/reports/saved/{report}/", "admin"),
    ("saved report run", "/reports/saved/{report}/run/", "admin"),
    ("expense", "/expenses/{expense}/", "admin"),
    ("gallery upload", "/gallery/{post}/", "member"),
    ("check-in", "/attendance/{checkin}/", "member"),
    ("referral", "/referrals/{referral}/", "member"),
    ("enquiry", "/crm/enquiries/{enquiry}/", "admin"),
]

#: State-changing actions, each tried against the target gym's row.
ACTIONS = [
    ("cancel a pt session", "/pt/sessions/{pt}/cancel/", {}),
    ("close off a pt session", "/pt/sessions/{pt}/complete/", {"is_paid": True}),
    ("book a class", "/schedule/{klass}/book/", {}),
    ("cancel a class booking", "/schedule/{klass}/cancel/", {}),
    ("approve an upload", "/gallery/{post}/approve/", {}),
    ("reward a referral", "/referrals/{referral}/reward/", {}),
    ("activate a diet plan", "/nutrition/plans/{diet}/activate/", {}),
    ("activate a split", "/workouts/splits/{split}/activate/", {}),
    ("convert an enquiry", "/crm/enquiries/{enquiry}/convert/", {}),
    ("reset a member's authenticator", "/auth/admin/users/{member}/reset-mfa/", {}),
    ("set a member's password", "/auth/admin/members/{member}/set-password/", {"password": "a-long-new-passphrase-77"}),
]

#: (list path, the target it must never contain, who at the target gym sees it)
LISTS = [
    ("/auth/admin/users/", "member", "admin"),
    ("/billing/admin/members/", "member", "admin"),
    ("/billing/admin/payments/", "payment", "admin"),
    ("/billing/admin/payments/?member={member}", "payment", "admin"),
    ("/billing/plans/", "plan", "admin"),
    ("/invoices/", "invoice", "member"),
    ("/invoices/?member={member}", "invoice", "admin"),
    ("/workouts/sessions/", "session", "member"),
    ("/workouts/sessions/?member={member}", "session", "trainer"),
    ("/workouts/splits/", "split", "member"),
    ("/workouts/splits/?member={member}", "split", "trainer"),
    ("/nutrition/plans/", "diet", "member"),
    ("/nutrition/plans/?member={member}", "diet", "trainer"),
    ("/bodystats/measurements/", "measurement", "member"),
    ("/bodystats/measurements/?member={member}", "measurement", "trainer"),
    ("/bodystats/goals/?member={member}", "goal", "trainer"),
    ("/pt/sessions/", "pt", "member"),
    ("/schedule/", "klass", "member"),
    ("/reports/saved/", "report", "admin"),
    ("/expenses/", "expense", "admin"),
    ("/gallery/", "post", "member"),
    ("/attendance/", "checkin", "member"),
    ("/referrals/", "referral", "member"),
    ("/referrals/?referrer={member}", "referral", "admin"),
    ("/crm/enquiries/", "enquiry", "admin"),
]


def gym_rows(gym, member, trainer, tag, price):
    """One of everything a gym owns, belonging to `member`, trained by `trainer`."""
    today = timezone.localdate()
    member.profile.trainer = trainer
    member.profile.save(update_fields=["trainer"])
    with context.scope(gym):
        plan = Plan.objects.create(name=f"Plan {tag}", price=Decimal(price), duration_days=30)
        payment = record_payment(
            member=member, plan=plan, amount=plan.price, method=PaymentMethod.CASH
        )
        invoice = issue_invoice(payment)
        exercise, _ = Exercise.objects.get_or_create(name="Matrix Squat")
        session = WorkoutSession.objects.create(user=member)
        log = WorkoutLog.objects.create(session=session, exercise=exercise, set_number=1, reps=5)
        split = WorkoutSplit.objects.create(user=member)
        split_day = SplitDay.objects.create(split=split, weekday=0)
        diet = DietPlan.objects.create(user=member)
        diet_day = DietDay.objects.create(plan=diet, weekday=0)
        meal = DietMeal.objects.create(day=diet_day, meal_type="lunch")
        measurement = BodyMeasurement.objects.create(user=member, weight_kg=Decimal("80"))
        goal = MemberGoal.objects.create(user=member, goal_type="weight", target_value=Decimal("70"))
        pt = PTSession.objects.create(
            trainer=trainer, member=member, date=today + timedelta(days=3),
            start_time=time(9), end_time=time(10),
        )
        klass = ClassSession.objects.create(
            title=f"Spin {tag}", trainer=trainer, date=today + timedelta(days=3),
            start_time=time(18), end_time=time(19),
        )
        ClassBooking.objects.create(member=member, session=klass)
        report = SavedReport.objects.create(
            name=f"Report {tag}", definition={"source": "payments", "columns": ["amount"]}
        )
        category = ExpenseCategory.objects.create(name=f"Rent {tag}", description="")
        expense = Expense.objects.create(category=category, amount=Decimal("500"))
        post = GalleryPost.objects.create(
            uploader=member, media=f"gallery/{tag}.png", media_type="image", approved=False
        )
        checkin = CheckInOut.objects.create(user=member)
        referral = Referral.objects.create(referrer=member, name=f"Friend {tag}")
        enquiry = Enquiry.objects.create(
            name=f"Lead {tag}", phone="+91 90000 00000", follow_up_on=today
        )
    return {
        "member": member, "payment": payment, "plan": plan, "invoice": invoice,
        "session": session, "log": log, "split": split, "split_day": split_day,
        "diet": diet, "diet_day": diet_day, "meal": meal, "measurement": measurement,
        "goal": goal, "pt": pt, "klass": klass, "report": report, "expense": expense,
        "category": category, "post": post, "checkin": checkin, "referral": referral,
        "enquiry": enquiry,
    }


def snapshot(rows):
    """Every target row as stored, so a refused change can be shown to be no change."""
    state = {}
    for key, obj in rows.items():
        model = type(obj)
        manager = getattr(model, "unscoped", model._default_manager)
        state[key] = manager.filter(pk=obj.pk).values().first()
    return state


def ids_in(resp):
    data = getattr(resp, "data", None)
    if isinstance(data, dict):
        data = data.get("results", data)
    if not isinstance(data, list):
        return set()
    return {row.get("id") for row in data if isinstance(row, dict)}


class CrossGymMatrix(TwoGymsTestCase):
    def setUp(self):
        super().setUp()
        self.admin_a = make_person("admin_a", gym=self.gym_a, role=Role.ADMIN)
        self.trainer_a = make_person("trainer_a", gym=self.gym_a, role=Role.TRAINER)
        self.member_a = make_person("member_a", gym=self.gym_a, first_name="Alma")
        self.member_a2 = make_person("member_a2", gym=self.gym_a, first_name="Abel")
        self.admin_b = make_person("admin_b", gym=self.gym_b, role=Role.ADMIN)
        self.trainer_b = make_person("trainer_b", gym=self.gym_b, role=Role.TRAINER)
        self.member_b = make_person("member_b", gym=self.gym_b, first_name="Beatrix", last_name="Outsider")

        self.rows_a = gym_rows(self.gym_a, self.member_a, self.trainer_a, "a", "1000")
        self.rows_b = gym_rows(self.gym_b, self.member_b, self.trainer_b, "b", "7777")
        self.people_a = {"admin": self.admin_a, "trainer": self.trainer_a, "member": self.member_a}
        self.people_b = {"admin": self.admin_b, "trainer": self.trainer_b, "member": self.member_b}

    def as_user(self, user):
        self.client.credentials()
        self.sign_in(user)

    def path(self, template, rows):
        return template.format(**{key: obj.pk for key, obj in rows.items()})

    def call(self, method, gym, path, data=None):
        url = self.at(gym) + path
        return getattr(self.client, method)(url, data or {}, format="json")


class TargetsAreRealTests(CrossGymMatrix):
    """The controls: gym B's own people reach every target, so refusals mean something."""

    def test_each_object_route_answers_its_entitled_reader(self):
        for label, template, role in OBJECT_ROUTES:
            with self.subTest(route=label):
                self.as_user(self.people_b[role])
                resp = self.call("get", self.gym_b, self.path(template, self.rows_b))
                self.assertEqual(resp.status_code, 200, f"{label}: {resp.status_code} {getattr(resp, 'data', '')}")

    def test_each_list_shows_the_target_to_its_entitled_reader(self):
        for template, key, role in LISTS:
            with self.subTest(route=template):
                self.as_user(self.people_b[role])
                resp = self.call("get", self.gym_b, self.path(template, self.rows_b))
                self.assertEqual(resp.status_code, 200, f"{template}: {resp.status_code}")
                if key != "member":
                    self.assertIn(self.rows_b[key].pk, ids_in(resp), template)


class CrossGymObjectTests(CrossGymMatrix):
    """Gym A's admin, trainer and member against gym B's rows, by id."""

    def test_nobody_at_gym_a_can_read_gym_b_rows_through_either_gyms_url(self):
        for role, actor in self.people_a.items():
            self.as_user(actor)
            for label, template, _ in OBJECT_ROUTES:
                for gym in (self.gym_a, self.gym_b):
                    with self.subTest(actor=role, route=label, via=gym.slug):
                        resp = self.call("get", gym, self.path(template, self.rows_b))
                        self.assertIn(resp.status_code, REFUSED, resp.content[:200])

    def test_nobody_at_gym_a_can_change_or_delete_gym_b_rows(self):
        before = snapshot(self.rows_b)
        changes = {"notes": "pwned", "weight_kg": "1", "name": "pwned", "price": "1", "amount": "1",
                   "title": "pwned", "status": "completed", "is_paid": True, "approved": True}
        for role, actor in self.people_a.items():
            self.as_user(actor)
            for label, template, _ in OBJECT_ROUTES:
                for gym in (self.gym_a, self.gym_b):
                    for method in ("patch", "put", "delete"):
                        with self.subTest(actor=role, route=label, via=gym.slug, method=method):
                            resp = self.call(method, gym, self.path(template, self.rows_b), changes)
                            self.assertIn(resp.status_code, REFUSED_WRITE, resp.content[:200])
        self.assertEqual(snapshot(self.rows_b), before)

    def test_nobody_at_gym_a_can_act_on_gym_b_rows(self):
        before = snapshot(self.rows_b)
        bookings = ClassBooking.unscoped.filter(session=self.rows_b["klass"]).count()
        for role, actor in self.people_a.items():
            self.as_user(actor)
            for label, template, body in ACTIONS:
                for gym in (self.gym_a, self.gym_b):
                    with self.subTest(actor=role, action=label, via=gym.slug):
                        resp = self.call("post", gym, self.path(template, self.rows_b), body)
                        self.assertIn(resp.status_code, REFUSED_WRITE, resp.content[:200])
        self.assertEqual(snapshot(self.rows_b), before)
        self.assertEqual(ClassBooking.unscoped.filter(session=self.rows_b["klass"]).count(), bookings)

    def test_gym_b_admins_cannot_reach_gym_a_either(self):
        """The same wall, the other way round, for the admin who can do the most."""
        self.as_user(self.admin_b)
        for label, template, _ in OBJECT_ROUTES:
            for gym in (self.gym_a, self.gym_b):
                with self.subTest(route=label, via=gym.slug):
                    resp = self.call("get", gym, self.path(template, self.rows_a))
                    self.assertIn(resp.status_code, REFUSED, resp.content[:200])


class CrossGymListTests(CrossGymMatrix):
    """Lists, including their `?member=` filters pointed at the other gym's people."""

    def test_no_list_at_gym_a_contains_gym_b_rows(self):
        for role, actor in self.people_a.items():
            self.as_user(actor)
            for template, key, _ in LISTS:
                with self.subTest(actor=role, route=template):
                    resp = self.call("get", self.gym_a, self.path(template, self.rows_b))
                    if resp.status_code != 200:
                        self.assertIn(resp.status_code, REFUSED | {400})
                        continue
                    self.assertNotIn(self.rows_b[key].pk, ids_in(resp))
                    self.assertNotIn(b"member_b", resp.content)
                    self.assertNotIn(b"Outsider", resp.content)

    def test_gym_b_lists_are_closed_to_gym_a_people(self):
        for role, actor in self.people_a.items():
            self.as_user(actor)
            for template, key, _ in LISTS:
                with self.subTest(actor=role, route=template):
                    resp = self.call("get", self.gym_b, self.path(template, self.rows_b))
                    self.assertIn(resp.status_code, REFUSED, resp.content[:200])


class CrossGymStrangerWriteTests(CrossGymMatrix):
    """Filing new rows under a gym you do not belong to."""

    def count_at_b(self):
        models = (WorkoutSession, BodyMeasurement, DietPlan, PTSession, Referral, GalleryPost,
                  CheckInOut, Expense, SavedReport, Payment)
        counts = {model.__name__: model.unscoped.filter(tenant=self.gym_b).count() for model in models}
        counts["Plan"] = Plan.unscoped.filter(organisation=self.gym_b.organisation).count()
        return counts

    def test_nobody_at_gym_a_can_create_rows_at_gym_b(self):
        day = (timezone.localdate() + timedelta(days=5)).isoformat()
        writes = [
            ("/workouts/sessions/", {}, "json"),
            ("/bodystats/measurements/", {"weight_kg": "70"}, "json"),
            ("/nutrition/plans/", {"name": "planted"}, "json"),
            ("/pt/sessions/", {"trainer": self.trainer_b.pk, "member": self.member_b.pk, "date": day,
                               "start_time": "11:00", "end_time": "12:00"}, "json"),
            ("/referrals/", {"name": "planted"}, "json"),
            ("/attendance/check_in/", {}, "json"),
            ("/expenses/", {"category": self.rows_b["category"].pk, "amount": "1"}, "json"),
            ("/reports/saved/", {"name": "planted", "definition": {"source": "payments", "columns": ["amount"]}}, "json"),
            ("/billing/plans/", {"name": "planted", "price": "1", "duration_days": 1}, "json"),
            ("/billing/admin/payments/", {"member": self.member_b.pk, "plan": self.rows_b["plan"].pk,
                                          "amount": "1", "method": "cash"}, "json"),
            ("/billing/checkout/", {"member": self.member_b.pk, "plan": self.rows_b["plan"].pk}, "json"),
        ]
        before = self.count_at_b()
        for role, actor in self.people_a.items():
            self.as_user(actor)
            for path, body, fmt in writes:
                with self.subTest(actor=role, route=path):
                    resp = self.client.post(self.at(self.gym_b) + path, body, format=fmt)
                    self.assertIn(resp.status_code, REFUSED_STRANGER, resp.content[:200])
            with self.subTest(actor=role, route="/gallery/ upload"):
                upload = SimpleUploadedFile("p.png", png_bytes(), content_type="image/png")
                resp = self.client.post(
                    self.at(self.gym_b) + "/gallery/", {"media": upload, "media_type": "image"},
                    format="multipart",
                )
                self.assertIn(resp.status_code, REFUSED_STRANGER, resp.content[:200])
        self.assertEqual(self.count_at_b(), before)

    def test_a_gym_a_admin_cannot_sell_to_a_gym_b_member_through_their_own_gym(self):
        self.as_user(self.admin_a)
        payments = Payment.unscoped.filter(member=self.member_b).count()
        for path in ("/billing/admin/payments/", "/billing/checkout/", "/billing/checkout/quote/"):
            with self.subTest(route=path):
                resp = self.call("post", self.gym_a, path, {
                    "member": self.member_b.pk, "plan": self.rows_a["plan"].pk,
                    "amount": "1000", "method": "cash",
                })
                self.assertIn(resp.status_code, {400, 403, 404}, resp.content[:200])
        self.assertEqual(Payment.unscoped.filter(member=self.member_b).count(), payments)


class SameGymAuthorizationTests(CrossGymMatrix):
    """Inside one gym: other members' rows, unassigned members, and admin-only routes."""

    MEMBER_PRIVATE = [
        ("invoice", "/invoices/{invoice}/"),
        ("invoice pdf", "/invoices/{invoice}/pdf/"),
        ("workout session", "/workouts/sessions/{session}/"),
        ("workout log", "/workouts/logs/{log}/"),
        ("split", "/workouts/splits/{split}/"),
        ("split day", "/workouts/split-days/{split_day}/"),
        ("diet plan", "/nutrition/plans/{diet}/"),
        ("diet day", "/nutrition/days/{diet_day}/"),
        ("diet meal", "/nutrition/meals/{meal}/"),
        ("measurement", "/bodystats/measurements/{measurement}/"),
        ("goal", "/bodystats/goals/{goal}/"),
        ("pt session", "/pt/sessions/{pt}/"),
        ("check-in", "/attendance/{checkin}/"),
        ("referral", "/referrals/{referral}/"),
        ("unapproved upload", "/gallery/{post}/"),
    ]

    def test_a_member_cannot_read_change_or_delete_another_members_rows(self):
        before = snapshot(self.rows_a)
        self.as_user(self.member_a2)
        for label, template in self.MEMBER_PRIVATE:
            for method in ("get", "patch", "delete"):
                with self.subTest(route=label, method=method):
                    resp = self.call(method, self.gym_a, self.path(template, self.rows_a), {"notes": "pwned"})
                    allowed = REFUSED if method == "get" else REFUSED_WRITE
                    self.assertIn(resp.status_code, allowed, resp.content[:200])
        self.assertEqual(snapshot(self.rows_a), before)

    def test_a_member_cannot_reach_another_member_through_list_filters(self):
        self.as_user(self.member_a2)
        for template, key, _ in LISTS:
            if "{member}" not in template:
                continue
            with self.subTest(route=template):
                resp = self.call("get", self.gym_a, self.path(template, self.rows_a))
                if resp.status_code == 200:
                    self.assertNotIn(self.rows_a[key].pk, ids_in(resp))
                else:
                    self.assertIn(resp.status_code, REFUSED | {400})

    def test_a_trainer_cannot_reach_a_member_who_is_not_theirs(self):
        other_trainer = make_person("trainer_a2", gym=self.gym_a, role=Role.TRAINER)
        self.as_user(other_trainer)
        resp = self.call("get", self.gym_a, self.path("/auth/trainer/members/{member}/", self.rows_a))
        self.assertIn(resp.status_code, REFUSED)
        for template, key, _ in LISTS:
            if "{member}" not in template:
                continue
            with self.subTest(route=template):
                resp = self.call("get", self.gym_a, self.path(template, self.rows_a))
                if resp.status_code == 200:
                    self.assertNotIn(self.rows_a[key].pk, ids_in(resp))
                else:
                    self.assertIn(resp.status_code, REFUSED | {400})

    def test_members_and_trainers_cannot_use_admin_routes(self):
        admin_only = [
            "/billing/admin/payments/{payment}/", "/reports/saved/{report}/", "/expenses/{expense}/",
            "/crm/enquiries/{enquiry}/", "/auth/admin/users/{member}/", "/billing/admin/members/",
            "/auth/admin/members/export/", "/billing/admin/members/export/",
        ]
        for role in ("member", "trainer"):
            self.as_user(self.people_a[role])
            for template in admin_only:
                with self.subTest(actor=role, route=template):
                    resp = self.call("get", self.gym_a, self.path(template, self.rows_a))
                    self.assertIn(resp.status_code, REFUSED, resp.content[:200])


class MalformedAndUnknownIdTests(CrossGymMatrix):
    """An id that is not an id, or not anybody's, is an answer -- never a crash."""

    def test_no_object_route_answers_a_bad_id_with_a_server_error(self):
        placeholders = dict.fromkeys(self.rows_a, None)
        for role in ("admin", "member"):
            self.as_user(self.people_a[role])
            for label, template, _ in OBJECT_ROUTES + [(a, t, None) for a, t, _ in ACTIONS]:
                for bad in ("abc", "-5", "99999999", "1.5", "%20"):
                    path = template.format(**{key: bad for key in placeholders})
                    method = "post" if template.endswith(("/cancel/", "/complete/", "/book/", "/approve/",
                                                          "/reward/", "/activate/", "/convert/",
                                                          "/reset-mfa/", "/set-password/")) else "get"
                    with self.subTest(actor=role, route=label, id=bad):
                        resp = self.call(method, self.gym_a, path)
                        self.assertLess(resp.status_code, 500, resp.content[:200])
                        self.assertIn(resp.status_code, {400, 403, 404, 405})


class CrossGymExportAndReminderTests(CrossGymMatrix):
    """What leaves the system in bulk: spreadsheets and reminder emails."""

    def cells(self, resp):
        sheet = openpyxl.load_workbook(io.BytesIO(resp.content)).active
        return " ".join(str(cell) for row in sheet.iter_rows(values_only=True) for cell in row if cell is not None)

    def test_gym_a_exports_name_nobody_from_gym_b(self):
        self.as_user(self.admin_a)
        for path in ("/auth/admin/members/export/", "/billing/admin/members/export/"):
            with self.subTest(export=path):
                resp = self.call("get", self.gym_a, path)
                self.assertEqual(resp.status_code, 200)
                text = self.cells(resp)
                self.assertIn("member_a", text)
                self.assertNotIn("member_b", text)
                self.assertNotIn("Outsider", text)

    def test_a_gym_a_report_export_carries_no_gym_b_money(self):
        self.as_user(self.admin_a)
        resp = self.client.post(
            self.at(self.gym_a) + "/reports/export/",
            {"source": "payments", "columns": ["amount"]}, format="json",
        )
        self.assertEqual(resp.status_code, 200, resp.content[:200])
        text = self.cells(resp)
        self.assertIn("1000", text)
        self.assertNotIn("7777", text)

    def test_a_gym_b_saved_report_cannot_be_run_or_exported_from_gym_a(self):
        self.as_user(self.admin_a)
        resp = self.call("get", self.gym_a, self.path("/reports/saved/{report}/run/", self.rows_b))
        self.assertIn(resp.status_code, REFUSED)

    def test_reminders_at_gym_a_never_reach_gym_b_members(self):
        period_end = Payment.unscoped.get(pk=self.rows_b["payment"].pk).period_end
        for window in EXPIRY_WINDOWS:
            on = period_end - timedelta(days=window)
            with self.subTest(window=window):
                with context.scope(self.gym_a):
                    at_a = [member for member, _, _ in expiring_members(on=on)]
                with context.scope(self.gym_b):
                    at_b = [member for member, _, _ in expiring_members(on=on)]
                self.assertNotIn(self.member_b, at_a)
                self.assertIn(self.member_b, at_b)
        self.assertEqual(len(mail.outbox), 0)


GATEWAY = {
    "RAZORPAY_KEY_ID": "rzp_test_key",
    "RAZORPAY_KEY_SECRET": "rzp_test_secret",
    "RAZORPAY_WEBHOOK_SECRET": "hook_secret",
}


def gateway_order(order_id="order_MATRIX"):
    return {"id": order_id, "amount": 777700, "currency": "INR", "status": "created"}


@override_settings(**GATEWAY)
class CrossGymOnlineCheckoutTests(CrossGymMatrix):
    """Paying online at a gym you do not belong to."""

    @patch("billing.gateway.create_order")
    def test_nobody_at_gym_a_can_open_an_order_at_gym_b(self, create_order):
        create_order.return_value = gateway_order()
        before = PaymentOrder.unscoped.filter(tenant=self.gym_b).count()
        for role, actor in self.people_a.items():
            self.as_user(actor)
            with self.subTest(actor=role):
                resp = self.client.post(
                    self.at(self.gym_b) + "/billing/online/order/",
                    {"plan": self.rows_b["plan"].pk}, format="json",
                )
                self.assertIn(resp.status_code, REFUSED_STRANGER, resp.content[:200])
        self.assertEqual(PaymentOrder.unscoped.filter(tenant=self.gym_b).count(), before)
        create_order.assert_not_called()

    @patch("billing.gateway.create_order")
    def test_a_gym_b_member_still_opens_an_order_at_their_own_gym(self, create_order):
        create_order.return_value = gateway_order("order_OWN")
        self.as_user(self.member_b)
        resp = self.client.post(
            self.at(self.gym_b) + "/billing/online/order/",
            {"plan": self.rows_b["plan"].pk}, format="json",
        )
        self.assertEqual(resp.status_code, 201, resp.content[:200])

    def test_payment_config_and_verification_are_closed_to_other_gyms(self):
        self.as_user(self.member_a)
        config = self.client.get(self.at(self.gym_b) + "/billing/online/config/")
        self.assertIn(config.status_code, REFUSED_STRANGER, config.content[:200])
        verify = self.client.post(
            self.at(self.gym_b) + "/billing/online/verify/",
            {"razorpay_order_id": "order_x", "razorpay_payment_id": "pay_x", "razorpay_signature": "sig"},
            format="json",
        )
        self.assertIn(verify.status_code, REFUSED_STRANGER | {400}, verify.content[:200])

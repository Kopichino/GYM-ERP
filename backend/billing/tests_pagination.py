"""Every billing record is reachable, one page at a time, and only this gym's.

The billing lists are paginated at twenty rows. The Billing pages read the first
page and treated it as the whole list -- so a gym with 22 payments saw 20, and
nothing said there was more. Invoices were fetched as a separate list and
matched to payments by id, so a payment whose invoice fell on the second page
offered to issue one that already existed.

These pin the API side of the fix: pages that add up to the whole set, a page
size the client may raise within a cap, the member filter holding on every page,
each payment carrying its own invoice, and nothing from another gym on any page.
"""

import io
from datetime import timedelta
from decimal import Decimal

import openpyxl
from django.contrib.auth import get_user_model
from django.utils import timezone
from rest_framework.test import APITestCase

from accounts.models import MemberProfile, Role
from core.testing import TenantAPIMixin, founding_tenant
from invoicing.models import Invoice
from invoicing.services import issue_invoice
from tenancy import context
from tenancy.models import Membership

from .models import Payment, PaymentMethod, PaymentStatus, Plan
from .services import record_payment

User = get_user_model()

TODAY = timezone.localdate()


class BillingPaginationTests(TenantAPIMixin, APITestCase):
    def setUp(self):
        self.admin = self._person("owner", Role.ADMIN)
        self.plan = Plan.objects.create(name="Monthly", price=Decimal("1000"), duration_days=30)
        self.client.force_authenticate(self.admin)

    def _person(self, username, role=Role.MEMBER):
        user = User.objects.create_user(
            username=username, email=f"{username}@example.com", password="pass12345", role=role
        )
        MemberProfile.objects.get_or_create(user=user)
        self.member_for(user, role)
        return user

    def _pay(self, member, days_ago=0, invoice=True, plan=None):
        payment = record_payment(
            member=member,
            plan=plan or self.plan,
            amount=Decimal("1000"),
            method=PaymentMethod.CASH,
            paid_date=TODAY - timedelta(days=days_ago),
        )
        if invoice:
            issue_invoice(payment)
        return payment

    def _all_pages(self, path, **params):
        rows, page = [], 1
        while True:
            resp = self.client.get(path, {**params, "page": page})
            self.assertEqual(resp.status_code, 200, resp.content)
            rows.extend(resp.data["results"])
            if not resp.data["next"]:
                return resp.data["count"], rows
            page += 1

    # -- page boundaries

    def test_a_single_payment_is_one_page(self):
        member = self._person("solo")
        payment = self._pay(member)
        resp = self.client.get("/api/billing/admin/payments/")
        self.assertEqual((resp.data["count"], resp.data["next"]), (1, None))
        self.assertEqual([row["id"] for row in resp.data["results"]], [payment.pk])

    def test_exactly_twenty_payments_fit_on_one_page(self):
        member = self._person("twenty")
        for day in range(20):
            self._pay(member, days_ago=day, invoice=False)
        resp = self.client.get("/api/billing/admin/payments/")
        self.assertEqual(resp.data["count"], 20)
        self.assertEqual(len(resp.data["results"]), 20)
        self.assertIsNone(resp.data["next"])

    def test_the_twenty_first_payment_is_on_a_second_page_not_lost(self):
        member = self._person("spill")
        made = {self._pay(member, days_ago=day, invoice=False).pk for day in range(21)}
        first = self.client.get("/api/billing/admin/payments/")
        self.assertEqual(first.data["count"], 21)
        self.assertEqual(len(first.data["results"]), 20)
        self.assertIsNotNone(first.data["next"])
        second = self.client.get("/api/billing/admin/payments/", {"page": 2})
        self.assertEqual(len(second.data["results"]), 1)
        seen = [row["id"] for row in first.data["results"] + second.data["results"]]
        self.assertEqual(len(seen), len(set(seen)))
        self.assertEqual(set(seen), made)

    def test_every_payment_is_reachable_across_several_pages(self):
        member = self._person("many")
        made = {self._pay(member, days_ago=day, invoice=False).pk for day in range(45)}
        count, rows = self._all_pages("/api/billing/admin/payments/")
        self.assertEqual(count, 45)
        self.assertEqual(sorted(row["id"] for row in rows), sorted(made))

    def test_the_page_size_can_be_raised_but_not_without_limit(self):
        member = self._person("bulk")
        Payment.objects.bulk_create(
            [
                Payment(
                    tenant=self.tenant,
                    member=member,
                    plan=self.plan,
                    amount=Decimal("10"),
                    method=PaymentMethod.CASH,
                    status=PaymentStatus.COMPLETED,
                    paid_date=TODAY,
                    period_start=TODAY,
                    period_end=TODAY + timedelta(days=30),
                )
                for _ in range(105)
            ]
        )
        self.assertEqual(
            len(self.client.get("/api/billing/admin/payments/", {"page_size": 50}).data["results"]),
            50,
        )
        capped = self.client.get("/api/billing/admin/payments/", {"page_size": 1000}).data
        self.assertEqual(len(capped["results"]), 100)
        self.assertEqual(capped["count"], 105)

    # -- filters and invoices hold on every page

    def test_the_member_filter_holds_on_every_page(self):
        regular, other = self._person("regular"), self._person("other")
        for day in range(25):
            self._pay(regular, days_ago=day, invoice=False)
        for day in range(5):
            self._pay(other, days_ago=day, invoice=False)
        count, rows = self._all_pages("/api/billing/admin/payments/", member=regular.pk)
        self.assertEqual(count, 25)
        self.assertEqual({row["member"] for row in rows}, {regular.pk})
        second = self.client.get("/api/billing/admin/payments/", {"member": regular.pk, "page": 2})
        self.assertEqual(len(second.data["results"]), 5)

    def test_each_payment_carries_its_own_invoice_on_every_page(self):
        member = self._person("invoiced")
        for day in range(22):
            self._pay(member, days_ago=day)
        count, rows = self._all_pages("/api/billing/admin/payments/")
        self.assertEqual(count, 22)
        invoices = {inv.payment_id: inv for inv in Invoice.objects.all()}
        self.assertEqual(len(invoices), count, "payments and invoices disagree on the total")
        for row in rows:
            with self.subTest(payment=row["id"]):
                self.assertEqual(row["invoice"]["id"], invoices[row["id"]].pk)
                self.assertEqual(row["invoice"]["number"], invoices[row["id"]].number)

    def test_a_payment_without_an_invoice_says_so(self):
        self._pay(self._person("uninvoiced"), invoice=False)
        resp = self.client.get("/api/billing/admin/payments/")
        self.assertIsNone(resp.data["results"][0]["invoice"])

    # -- the member billing list and its export

    def test_the_billing_member_list_counts_and_pages_through_every_member(self):
        names = {self._person(f"member{n:02d}").username for n in range(23)}
        count, rows = self._all_pages("/api/billing/admin/members/")
        self.assertEqual(count, 23)
        self.assertEqual({row["username"] for row in rows}, names)
        second = self.client.get("/api/billing/admin/members/", {"page": 2})
        self.assertEqual(len(second.data["results"]), 3)

    def test_the_billing_export_includes_every_member_not_just_the_first_page(self):
        names = {self._person(f"member{n:02d}").username for n in range(25)}
        resp = self.client.get("/api/billing/admin/members/export/")
        self.assertEqual(resp.status_code, 200)
        sheet = openpyxl.load_workbook(io.BytesIO(resp.content)).active
        exported = {row[0] for row in sheet.iter_rows(min_row=2, values_only=True)}
        self.assertEqual(exported, names)

    def test_the_member_list_behind_the_till_can_be_read_in_full(self):
        names = {self._person(f"member{n:02d}").username for n in range(25)}
        resp = self.client.get("/api/auth/admin/members/", {"page_size": 100})
        self.assertEqual(resp.data["count"], 25)
        self.assertEqual({row["username"] for row in resp.data["results"]}, names)

    # -- a member's own billing

    def test_a_member_can_reach_every_payment_and_invoice_of_their_own(self):
        member = self._person("longstanding")
        for month in range(23):
            self._pay(member, days_ago=30 * month)
        self.client.force_authenticate(member)
        count, rows = self._all_pages("/api/billing/my-payments/")
        self.assertEqual(count, 23)
        self.assertTrue(all(row["invoice"] and row["invoice"]["number"] for row in rows))

    # -- another gym, on any page

    def test_no_page_includes_another_gyms_members_or_payments(self):
        local = self._person("local")
        for day in range(21):
            self._pay(local, days_ago=day, invoice=False)

        _, rival = founding_tenant("rivalgym")
        outsider = User.objects.create_user(
            username="outsider", email="outsider@example.com", password="pass12345"
        )
        MemberProfile.objects.get_or_create(user=outsider)
        Membership.objects.create(user=outsider, tenant=rival, role=Role.MEMBER)
        with context.scope(rival):
            rival_plan = Plan.objects.create(name="Monthly", price=Decimal("800"), duration_days=30)
            for day in range(21):
                self._pay(outsider, days_ago=day, plan=rival_plan, invoice=False)

        count, rows = self._all_pages("/api/billing/admin/payments/")
        self.assertEqual(count, 21)
        self.assertNotIn(outsider.pk, {row["member"] for row in rows})

        count, rows = self._all_pages("/api/billing/admin/members/")
        self.assertEqual({row["username"] for row in rows}, {"local"})

        resp = self.client.get("/api/billing/admin/members/export/")
        sheet = openpyxl.load_workbook(io.BytesIO(resp.content)).active
        self.assertNotIn("outsider", {row[0] for row in sheet.iter_rows(min_row=2, values_only=True)})

    def test_only_an_admin_reads_the_ledger(self):
        member = self._person("curious")
        self.client.force_authenticate(member)
        self.assertEqual(self.client.get("/api/billing/admin/payments/").status_code, 403)
        self.assertEqual(self.client.get("/api/billing/admin/members/").status_code, 403)
